"""Document exports — the full transcript and the editor's selects as
Word (.docx), PDF, or Excel (.xlsx).

These are the "Documents" exports on the project Export tab: the things an
editor hands to a producer, a client, or a transcript reviewer — not the
NLE timelines (those live in ``fcpxml_export`` / ``exporters.fcpxml`` /
``exporters.premiere_xml`` / ``exporters.resolve_xml`` and are untouched by
this module).

Generators:

- DOCX via python-docx, PDF via reportlab — both already in the app bundle
  (the quote-sheet extension uses them).
- XLSX via :func:`write_xlsx`, a dependency-free SpreadsheetML writer
  (zip + XML, inline strings, bold frozen header row, sized columns). No
  openpyxl in the bundle, and a flat table doesn't need one.

Text coverage: reportlab's built-in Helvetica is Latin-1 only, which turns
Polish/Czech/Turkish/Greek/Cyrillic transcripts into boxes. PDF output
registers macOS's Arial Unicode (Supplemental fonts, present on every Mac)
at runtime via :func:`register_unicode_fonts`; Helvetica stays the
fallback when the file is missing. Nothing is bundled.

Speaker labels resolve through the project's ``speaker_names`` rename map
so no raw ``SPEAKER_NN`` leaks into a deliverable.
"""

from __future__ import annotations

import os
import re
import zipfile
from datetime import datetime
from typing import Any, Iterable
from xml.sax.saxutils import escape as _xml_escape


# ── Fonts ────────────────────────────────────────────────────────────

_FONT_REGULAR_CANDIDATES = (
    '/System/Library/Fonts/Supplemental/Arial Unicode.ttf',
    '/Library/Fonts/Arial Unicode.ttf',
    '/System/Library/Fonts/Supplemental/Arial.ttf',
)
_FONT_BOLD_CANDIDATES = (
    '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
    '/Library/Fonts/Arial Bold.ttf',
)
_fonts: tuple[str, str] | None = None


def register_unicode_fonts() -> tuple[str, str]:
    """Register a Unicode TrueType family with reportlab (once per process).

    Returns ``(regular_font_name, bold_font_name)``. Falls back to
    ``('Helvetica', 'Helvetica-Bold')`` when no usable TTF is on disk or
    reportlab is unavailable, so callers can always use the names directly.
    The bold face is Arial Bold (Latin/Greek/Cyrillic); body text — the
    part that carries the transcript — uses Arial Unicode's full coverage.
    """
    global _fonts
    if _fonts is not None:
        return _fonts
    fallback = ('Helvetica', 'Helvetica-Bold')
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.fonts import addMapping
    except Exception:
        _fonts = fallback
        return _fonts
    regular = next((p for p in _FONT_REGULAR_CANDIDATES if os.path.isfile(p)), None)
    if regular is None:
        _fonts = fallback
        return _fonts
    bold = next((p for p in _FONT_BOLD_CANDIDATES if os.path.isfile(p)), regular)
    try:
        pdfmetrics.registerFont(TTFont('DozaSans', regular))
        pdfmetrics.registerFont(TTFont('DozaSans-Bold', bold))
        addMapping('DozaSans', 0, 0, 'DozaSans')
        addMapping('DozaSans', 1, 0, 'DozaSans-Bold')
        addMapping('DozaSans', 0, 1, 'DozaSans')
        addMapping('DozaSans', 1, 1, 'DozaSans-Bold')
        _fonts = ('DozaSans', 'DozaSans-Bold')
    except Exception:
        _fonts = fallback
    return _fonts


# ── Timecode helpers ────────────────────────────────────────────────

def fmt_clock(seconds: float) -> str:
    """``HH:MM:SS`` for a media-relative position."""
    try:
        s = max(0, int(float(seconds) + 0.5))
    except (TypeError, ValueError):
        s = 0
    return f'{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}'


def fmt_duration(seconds: float) -> str:
    """``1:23`` / ``12s`` style human duration."""
    try:
        s = max(0, int(float(seconds) + 0.5))
    except (TypeError, ValueError):
        s = 0
    if s < 60:
        return f'{s}s'
    return f'{s // 60}:{s % 60:02d}'


def fmt_source_tc(seconds: float, start_tc: dict | None) -> str:
    """``HH:MM:SS:FF`` in the media's embedded timecode, or ``''`` when the
    project carries no usable ``start_tc`` (the dashboard's lazy probe:
    ``{frames, fps, drop, raw}``). Non-drop arithmetic; drop-frame media
    gets the same frame count labelled as the NLE would display it only
    approximately, so we skip the column for drop-frame sources rather
    than show a wrong number."""
    if not isinstance(start_tc, dict):
        return ''
    try:
        fps = float(start_tc.get('fps') or 0)
        base = int(start_tc.get('frames') or 0)
    except (TypeError, ValueError):
        return ''
    if fps <= 0 or start_tc.get('drop'):
        return ''
    nominal = int(round(fps))
    total = base + int(round(float(seconds) * fps))
    ff = total % nominal
    secs = total // nominal
    return f'{secs // 3600:02d}:{(secs % 3600) // 60:02d}:{secs % 60:02d}:{ff:02d}'


# ── Data collection ──────────────────────────────────────────────────

def _speaker_resolver(project: dict):
    names = project.get('speaker_names') or {}

    def resolve(raw):
        raw = (raw or '').strip()
        if not raw:
            return ''
        mapped = names.get(raw)
        return mapped.strip() if isinstance(mapped, str) and mapped.strip() else raw
    return resolve


def transcript_rows(project: dict) -> list[dict[str, Any]]:
    """One row per transcript segment: start/end seconds, resolved speaker,
    text, and the source timecode when the media carries one."""
    transcript = project.get('transcript') or {}
    segments = transcript.get('segments') or []
    resolve = _speaker_resolver(project)
    start_tc = project.get('start_tc')
    rows = []
    for seg in segments:
        try:
            start = float(seg.get('start', 0) or 0)
            end = float(seg.get('end', start) or start)
        except (TypeError, ValueError):
            continue
        text = (seg.get('text') or '').strip()
        if not text:
            continue
        rows.append({
            'start': start,
            'end': end,
            'speaker': resolve(seg.get('speaker')),
            'text': text,
            'source_tc': fmt_source_tc(start, start_tc),
        })
    return rows


def transcript_paragraphs(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    """Merge consecutive same-speaker segments into readable paragraphs
    (the Word/PDF shape); the Excel sheet keeps per-segment rows."""
    out: list[dict[str, Any]] = []
    for r in rows:
        if out and out[-1]['speaker'] == r['speaker']:
            out[-1]['text'] += ' ' + r['text']
            out[-1]['end'] = r['end']
        else:
            out.append(dict(r))
    return out


_CATEGORY_TITLES = {
    'labels': 'My Clips',
    'social': 'Social Clips',
    'story': 'Story Beats',
    'soundbites': 'Soundbites',
}


def _to_seconds(val) -> float:
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val or '').strip()
    if ':' in s:
        parts = [float(p) for p in s.split(':')]
        if len(parts) == 3:
            return parts[0] * 3600 + parts[1] * 60 + parts[2]
        if len(parts) == 2:
            return parts[0] * 60 + parts[1]
    return float(s or 0)


def selects_rows(project: dict, categories: Iterable[str]) -> list[dict[str, Any]]:
    """The editor's selects, in the same category semantics as the NLE
    export (``_build_nle_export`` in app.py): ``labels`` = highlighted
    clips (My Clips, named by the color label), ``social`` = AI social
    clips, ``story`` = AI story beats, ``soundbites`` = strongest soundbites.
    Each row: category, label, start/end seconds, speaker, quote (the
    transcript words inside the range), note."""
    wanted = {c for c in (categories or []) if c in _CATEGORY_TITLES}
    resolve = _speaker_resolver(project)
    transcript = project.get('transcript') or {}
    segments = transcript.get('segments') or []
    analysis = project.get('analysis') or {}
    start_tc = project.get('start_tc')

    def words_in(start, end):
        bits = []
        for seg in segments:
            try:
                s0 = float(seg.get('start', 0) or 0)
                e0 = float(seg.get('end', s0) or s0)
            except (TypeError, ValueError):
                continue
            if e0 <= start or s0 >= end:
                continue
            words = seg.get('words') or []
            if words:
                for w in words:
                    try:
                        ws = float(w.get('start', 0) or 0)
                    except (TypeError, ValueError):
                        continue
                    if start <= ws < end:
                        bits.append((w.get('word') or '').strip())
            else:
                bits.append((seg.get('text') or '').strip())
        return ' '.join(b for b in bits if b)

    def speaker_at(start, end):
        for seg in segments:
            try:
                s0 = float(seg.get('start', 0) or 0)
                e0 = float(seg.get('end', s0) or s0)
            except (TypeError, ValueError):
                continue
            if s0 < end and e0 > start:
                return resolve(seg.get('speaker'))
        return ''

    rows: list[dict[str, Any]] = []

    def add(category, label, start, end, note):
        try:
            s = _to_seconds(start)
            e = _to_seconds(end)
        except ValueError:
            return
        if e <= s:
            e = s + 15
        rows.append({
            'category': category,
            'label': (label or '').strip(),
            'start': s,
            'end': e,
            'speaker': speaker_at(s, e),
            'quote': words_in(s, e),
            'note': (note or '').strip(),
            'source_tc': fmt_source_tc(s, start_tc),
            'source_tc_out': fmt_source_tc(e, start_tc),
        })

    if 'labels' in wanted:
        color_labels = project.get('color_labels') or {}
        for sec in project.get('labeled_sections') or []:
            name = color_labels.get(sec.get('color', ''), sec.get('color', '')) or 'Clip'
            add(name, name, sec.get('start', 0), sec.get('end', 0),
                sec.get('title') or sec.get('text'))
    if 'social' in wanted:
        for clip in analysis.get('social_clips') or []:
            add('Social Clip', clip.get('title'), clip.get('start', 0), clip.get('end', 0),
                clip.get('platform'))
    if 'story' in wanted:
        for beat in analysis.get('story_beats') or []:
            add('Story Beat', beat.get('label'), beat.get('start', 0),
                beat.get('end', beat.get('start', 0)), beat.get('description'))
    if 'soundbites' in wanted:
        for sb in analysis.get('strongest_soundbites') or []:
            add('Soundbite', (sb.get('text') or '')[:80], sb.get('start', 0),
                sb.get('end', sb.get('start', 0)), sb.get('why'))
    rows.sort(key=lambda r: r['start'])
    return rows


def project_meta_line(project: dict) -> str:
    bits = []
    src = project.get('filename') or os.path.basename(project.get('source_path') or '')
    if src:
        bits.append(src)
    transcript = project.get('transcript') or {}
    dur = transcript.get('duration')
    if dur:
        bits.append(f'{fmt_duration(dur)} long')
    lang = transcript.get('language')
    if lang:
        bits.append(f'language: {lang}')
    bits.append('exported ' + datetime.now().strftime('%Y-%m-%d %H:%M'))
    return ' · '.join(bits)


# ── XLSX writer (no dependencies) ────────────────────────────────────

_CONTROL_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _cell_xml(col_letter: str, row_idx: int, value, header: bool = False) -> str:
    ref = f'{col_letter}{row_idx}'
    style = ' s="1"' if header else ''
    if isinstance(value, bool):
        return f'<c r="{ref}" t="b"{style}><v>{1 if value else 0}</v></c>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{style}><v>{value}</v></c>'
    text = _CONTROL_CHARS.sub('', str(value if value is not None else ''))
    return (f'<c r="{ref}" t="inlineStr"{style}><is><t xml:space="preserve">'
            f'{_xml_escape(text)}</t></is></c>')


def _col_letter(n: int) -> str:
    s = ''
    n += 1
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def write_xlsx(path: str, sheets: list[tuple[str, list[str], list[list[Any]]]],
               widths: dict[str, list[int]] | None = None) -> str:
    """Write a minimal but well-formed .xlsx: one worksheet per
    ``(name, headers, rows)``; bold header row, frozen at row 1; column
    widths from ``widths[name]`` (character units) or a heuristic."""
    widths = widths or {}
    ws_xml = []
    for name, headers, rows in sheets:
        cols = len(headers)
        w = widths.get(name) or [
            min(80, max(10, max([len(str(h))] + [len(str(r[i])) if i < len(r) else 0
                                                   for r in rows[:200]]) + 2))
            for i, h in enumerate(headers)
        ]
        cols_xml = ''.join(
            f'<col min="{i + 1}" max="{i + 1}" width="{w[i]}" customWidth="1"/>'
            for i in range(cols))
        lines = [f'<row r="1">' + ''.join(
            _cell_xml(_col_letter(i), 1, h, header=True) for i, h in enumerate(headers)) + '</row>']
        for ri, row in enumerate(rows, start=2):
            lines.append(f'<row r="{ri}">' + ''.join(
                _cell_xml(_col_letter(i), ri, row[i] if i < len(row) else '')
                for i in range(cols)) + '</row>')
        ws_xml.append(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" '
            'activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
            f'<cols>{cols_xml}</cols><sheetData>{"".join(lines)}</sheetData></worksheet>')

    sheet_entries = ''.join(
        f'<sheet name="{_xml_escape(name[:31])}" sheetId="{i + 1}" r:id="rId{i + 1}"/>'
        for i, (name, _h, _r) in enumerate(sheets))
    workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                f'<sheets>{sheet_entries}</sheets></workbook>')
    wb_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               + ''.join(
                   f'<Relationship Id="rId{i + 1}" '
                   'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
                   f'Target="worksheets/sheet{i + 1}.xml"/>' for i in range(len(sheets)))
               + f'<Relationship Id="rId{len(sheets) + 1}" '
               'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
               'Target="styles.xml"/></Relationships>')
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
              '<font><b/><sz val="11"/><name val="Calibri"/></font></fonts>'
              '<fills count="2"><fill><patternFill patternType="none"/></fill>'
              '<fill><patternFill patternType="gray125"/></fill></fills>'
              '<borders count="1"><border/></borders>'
              '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
              '<cellXfs count="2"><xf xfId="0"/><xf fontId="1" xfId="0" applyFont="1">'
              '<alignment wrapText="0"/></xf></cellXfs></styleSheet>')
    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                     '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                     + ''.join(
                         f'<Override PartName="/xl/worksheets/sheet{i + 1}.xml" '
                         'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                         for i in range(len(sheets)))
                     + '</Types>')
    root_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                 'Target="xl/workbook.xml"/></Relationships>')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('_rels/.rels', root_rels)
        z.writestr('xl/workbook.xml', workbook)
        z.writestr('xl/_rels/workbook.xml.rels', wb_rels)
        z.writestr('xl/styles.xml', styles)
        for i, xml in enumerate(ws_xml):
            z.writestr(f'xl/worksheets/sheet{i + 1}.xml', xml)
    return path


# ── Builders: transcript ─────────────────────────────────────────────

def build_transcript_docx(project: dict, rows: list[dict], path: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.add_heading(project.get('name') or 'Transcript', level=1)
    meta = doc.add_paragraph(project_meta_line(project))
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x6B, 0x72)
    for para in transcript_paragraphs(rows):
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(0)
        tc = para['source_tc'] or fmt_clock(para['start'])
        r = head.add_run(tc)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x72)
        if para['speaker']:
            r2 = head.add_run('   ' + para['speaker'])
            r2.bold = True
            r2.font.size = Pt(10)
        body = doc.add_paragraph(para['text'])
        body.paragraph_format.space_after = Pt(4)
    doc.save(path)
    return path


def _pdf_styles():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    regular, bold = register_unicode_fonts()
    return {
        'title': ParagraphStyle('DocTitle', fontName=bold, fontSize=16, leading=20, spaceAfter=4),
        'meta': ParagraphStyle('DocMeta', fontName=regular, fontSize=8.5, leading=11,
                               textColor=HexColor('#6B6B72'), spaceAfter=12),
        'head': ParagraphStyle('ParaHead', fontName=bold, fontSize=9.5, leading=12,
                               spaceBefore=8, spaceAfter=1),
        'tc': ParagraphStyle('ParaTc', fontName=regular, fontSize=8.5, leading=11,
                             textColor=HexColor('#6B6B72')),
        'body': ParagraphStyle('ParaBody', fontName=regular, fontSize=10, leading=14, spaceAfter=4),
        'section': ParagraphStyle('Section', fontName=bold, fontSize=12, leading=15,
                                  spaceBefore=14, spaceAfter=4),
        'note': ParagraphStyle('Note', fontName=regular, fontSize=9, leading=12,
                               textColor=HexColor('#6B6B72'), spaceAfter=6),
    }


def _pdf_doc(path: str, title: str):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.8 * inch, title=title)


def _p(text: str, style):
    from reportlab.platypus import Paragraph
    return Paragraph(_xml_escape(text or ''), style)


def build_transcript_pdf(project: dict, rows: list[dict], path: str) -> str:
    styles = _pdf_styles()
    story = [_p(project.get('name') or 'Transcript', styles['title']),
             _p(project_meta_line(project), styles['meta'])]
    for para in transcript_paragraphs(rows):
        tc = para['source_tc'] or fmt_clock(para['start'])
        head = f'<font color="#6B6B72" size="8.5">{_xml_escape(tc)}</font>'
        if para['speaker']:
            head += '&nbsp;&nbsp;&nbsp;' + _xml_escape(para['speaker'])
        from reportlab.platypus import Paragraph
        story.append(Paragraph(head, styles['head']))
        story.append(_p(para['text'], styles['body']))
    _pdf_doc(path, project.get('name') or 'Transcript').build(story)
    return path


def build_transcript_xlsx(project: dict, rows: list[dict], path: str) -> str:
    has_tc = any(r['source_tc'] for r in rows)
    headers = ['Start', 'End', 'Speaker', 'Text'] + (['Source TC'] if has_tc else [])
    data = [[fmt_clock(r['start']), fmt_clock(r['end']), r['speaker'], r['text']]
            + ([r['source_tc']] if has_tc else []) for r in rows]
    widths = {'Transcript': [10, 10, 22, 90] + ([13] if has_tc else [])}
    return write_xlsx(path, [('Transcript', headers, data)], widths)


# ── Builders: selects ────────────────────────────────────────────────

def _select_headline(r: dict) -> str:
    tc_in = r['source_tc'] or fmt_clock(r['start'])
    tc_out = r.get('source_tc_out') or fmt_clock(r['end'])
    dur = fmt_duration(r['end'] - r['start'])
    head = f"{r['label'] or r['category']} — {tc_in} → {tc_out} ({dur})"
    if r['speaker']:
        head += f" · {r['speaker']}"
    return head


def _group_selects(rows: list[dict]) -> list[tuple[str, list[dict]]]:
    order: list[str] = []
    groups: dict[str, list[dict]] = {}
    for r in rows:
        key = r['category']
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(r)
    return [(k, groups[k]) for k in order]


def build_selects_docx(project: dict, rows: list[dict], path: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.add_heading(f"{project.get('name') or 'Project'} — Selects", level=1)
    meta = doc.add_paragraph(f'{len(rows)} selects · ' + project_meta_line(project))
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x6B, 0x72)
    for category, items in _group_selects(rows):
        doc.add_heading(f'{category} ({len(items)})', level=2)
        for r in items:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(8)
            head.paragraph_format.space_after = Pt(0)
            run = head.add_run(_select_headline(r))
            run.bold = True
            run.font.size = Pt(10)
            if r['quote']:
                q = doc.add_paragraph(f'“{r["quote"]}”')
                q.paragraph_format.space_after = Pt(2)
            if r['note']:
                n = doc.add_paragraph(r['note'])
                n.runs[0].italic = True
                n.runs[0].font.size = Pt(9)
                n.runs[0].font.color.rgb = RGBColor(0x6B, 0x6B, 0x72)
    doc.save(path)
    return path


def build_selects_pdf(project: dict, rows: list[dict], path: str) -> str:
    styles = _pdf_styles()
    name = project.get('name') or 'Project'
    story = [_p(f'{name} — Selects', styles['title']),
             _p(f'{len(rows)} selects · ' + project_meta_line(project), styles['meta'])]
    for category, items in _group_selects(rows):
        story.append(_p(f'{category} ({len(items)})', styles['section']))
        for r in items:
            story.append(_p(_select_headline(r), styles['head']))
            if r['quote']:
                story.append(_p(f'“{r["quote"]}”', styles['body']))
            if r['note']:
                story.append(_p(r['note'], styles['note']))
    _pdf_doc(path, f'{name} — Selects').build(story)
    return path


def build_selects_xlsx(project: dict, rows: list[dict], path: str) -> str:
    has_tc = any(r['source_tc'] for r in rows)
    headers = ['Category', 'Label', 'In', 'Out', 'Duration', 'Speaker', 'Quote', 'Note'] \
        + (['Source TC'] if has_tc else [])
    data = [[r['category'], r['label'], fmt_clock(r['start']), fmt_clock(r['end']),
             fmt_duration(r['end'] - r['start']), r['speaker'], r['quote'], r['note']]
            + ([r['source_tc']] if has_tc else []) for r in rows]
    widths = {'Selects': [14, 28, 10, 10, 10, 20, 80, 40] + ([13] if has_tc else [])}
    return write_xlsx(path, [('Selects', headers, data)], widths)


# ── Entry point ──────────────────────────────────────────────────────

DOCUMENT_KINDS = ('transcript', 'selects')
FORMATS = {
    'docx': ('Word', '.docx',
             'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
    'pdf': ('PDF', '.pdf', 'application/pdf'),
    'xlsx': ('Excel', '.xlsx',
             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
}

_BUILDERS = {
    ('transcript', 'docx'): build_transcript_docx,
    ('transcript', 'pdf'): build_transcript_pdf,
    ('transcript', 'xlsx'): build_transcript_xlsx,
    ('selects', 'docx'): build_selects_docx,
    ('selects', 'pdf'): build_selects_pdf,
    ('selects', 'xlsx'): build_selects_xlsx,
}


def safe_filename(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', ' ', (name or '').strip())
    s = re.sub(r'\s+', ' ', s).strip(' .')
    return s[:80] or 'Doza Project'


def allocate_path(exports_dir: str, project_name: str, what: str, fmt: str) -> str:
    """``<exports>/documents/<Project> – Transcript.docx`` — human names,
    uniquified with `` (2)`` rather than timestamps so the Finder reveal
    lands on something a producer recognises."""
    out_dir = os.path.join(exports_dir, 'documents')
    os.makedirs(out_dir, exist_ok=True)
    label = {'transcript': 'Transcript', 'selects': 'Selects'}[what]
    ext = FORMATS[fmt][1]
    base = f'{safe_filename(project_name)} – {label}'
    candidate = os.path.join(out_dir, base + ext)
    n = 2
    while os.path.exists(candidate):
        candidate = os.path.join(out_dir, f'{base} ({n}){ext}')
        n += 1
    return candidate


def export_document(project: dict, what: str, fmt: str, exports_dir: str,
                    categories: Iterable[str] = ()) -> tuple[str, int]:
    """Build the requested document into ``exports_dir``. Returns
    ``(path, row_count)``. Raises ``ValueError`` with a user-facing message
    when there is nothing to export."""
    if what not in DOCUMENT_KINDS:
        raise ValueError(f'Unknown document "{what}".')
    if fmt not in FORMATS:
        raise ValueError(f'Unsupported format "{fmt}". Use docx, pdf, or xlsx.')
    if what == 'transcript':
        rows = transcript_rows(project)
        if not rows:
            raise ValueError('Transcript not ready yet — nothing to export.')
    else:
        rows = selects_rows(project, categories)
        if not rows:
            raise ValueError('No selects in the checked categories — highlight clips in the '
                             'transcript or run AI Analysis first.')
    path = allocate_path(exports_dir, project.get('name') or 'Doza Project', what, fmt)
    _BUILDERS[(what, fmt)](project, rows, path)
    return path, len(rows)
